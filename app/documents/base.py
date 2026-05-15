"""Shared helpers for document generation."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Default placeholder logo shipped with the package
_DEFAULT_LOGO = Path(__file__).parent.parent / "static" / "logo_placeholder.png"


def resolve_logo(logo_path: Path | None) -> Path | None:
    """Return the logo path to embed, or None if unavailable."""
    if logo_path and logo_path.is_file():
        return logo_path
    if _DEFAULT_LOGO.is_file():
        return _DEFAULT_LOGO
    return None


def make_temp_file(suffix: str) -> Path:
    """Create an empty named temp file and return its Path."""
    fd, tmp = tempfile.mkstemp(suffix=suffix)
    import os

    os.close(fd)
    return Path(tmp)


def set_landscape(doc: Document) -> None:
    """Switch the document to landscape orientation."""

    section = doc.sections[0]
    section.orientation = 1  # WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width


def add_logo(doc: Document, logo_path: Path, max_w: float = 2.2, max_h: float = 1.1) -> None:
    from PIL import Image as PILImage

    img = PILImage.open(logo_path)
    w, h = img.size
    scale = min(max_w / w, max_h / h)
    para = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run()
    run.add_picture(str(logo_path), width=Inches(w * scale), height=Inches(h * scale))


def set_cell_bg(cell, hex_color: str) -> None:
    """Set the background fill colour of a table cell (hex without #)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def style_header_cell(cell, text: str, font_size: int = 10) -> None:
    """Bold, centred text in a table header cell."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cell.text = text
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.runs[0]
    run.bold = True
    run.font.size = Pt(font_size)


def style_body_cell(cell, text: str, font_size: int = 9) -> None:
    """Centred text in a table body cell."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cell.text = text
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if para.runs:
        para.runs[0].font.size = Pt(font_size)


def add_signature_block(doc: Document, left_label: str, right_label: str) -> None:
    """Add a two-column signature block at the bottom of the document."""
    doc.add_paragraph()  # spacing
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"

    labels = [left_label, right_label]
    for col_idx, label in enumerate(labels):
        # Row 0: label
        style_header_cell(table.rows[0].cells[col_idx], label, font_size=10)
        # Row 1: blank space for signature
        table.rows[1].cells[col_idx].text = ""
        # Row 2: "Name / Surname" and "Signature" lines (populated by caller)
        table.rows[2].cells[col_idx].text = ""

    # Give the signature rows some height
    from docx.oxml import OxmlElement

    for row in table.rows[1:]:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "800")
        trPr.append(trHeight)


def add_signature_table(
    doc: Document,
    left_label: str,
    right_label: str,
    left_name: str,
    right_name: str,
    date_label: str,
    sig_label: str,
    name_label: str,
) -> None:
    """Two-column signature block with name and date rows."""
    doc.add_paragraph()
    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = "Table Grid"

    # Row 0: role labels
    style_header_cell(tbl.rows[0].cells[0], left_label)
    style_header_cell(tbl.rows[0].cells[1], right_label)

    # Row 1: name/surname
    style_body_cell(tbl.rows[1].cells[0], name_label)
    style_body_cell(tbl.rows[1].cells[1], name_label)

    # Row 2: pre-filled names (known parties)
    style_body_cell(tbl.rows[2].cells[0], left_name)
    style_body_cell(tbl.rows[2].cells[1], right_name)

    # Row 3: date + signature blank
    style_body_cell(tbl.rows[3].cells[0], f"{date_label}: ___________    {sig_label}: ___________")
    style_body_cell(tbl.rows[3].cells[1], f"{date_label}: ___________    {sig_label}: ___________")

    # Height for rows 2-3
    from docx.oxml import OxmlElement

    for row in tbl.rows[2:]:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "600")
        trPr.append(trHeight)
