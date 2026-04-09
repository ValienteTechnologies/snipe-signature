"""Checkout form (asset delivery receipt) DOCX generation."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.documents.base import (
    add_logo,
    add_signature_table,
    make_temp_file,
    resolve_logo,
    set_cell_bg,
    set_landscape,
    style_body_cell,
    style_header_cell,
)
from app.i18n import Labels
from app.snipeit.models import AssetWithActivity, User


def build_checkout_docx(
    enriched: list[AssetWithActivity],
    user: User,
    labels: Labels,
    logo_path: Path | None = None,
    template: str = "default",
) -> Path:
    """Generate a checkout receipt DOCX and return the path to the temp file."""
    doc = Document()
    set_landscape(doc)

    resolved_logo = resolve_logo(logo_path)
    if resolved_logo:
        add_logo(doc, resolved_logo)
    else:
        doc.add_paragraph()

    # Title
    title_para = doc.add_paragraph(labels.doc_checkout_title)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.runs[0]
    title_run.bold = True
    title_run.font.size = Pt(14)

    # Subtitle: user name + generated timestamp
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    sub = doc.add_paragraph(f"{user.display_name}  —  {now_str}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(10)

    # Condition note
    note = doc.add_paragraph(labels.doc_condition_note)
    note.runs[0].font.size = Pt(9)
    note.runs[0].italic = True

    doc.add_paragraph()

    # Asset table
    columns = [
        labels.col_asset_tag,
        labels.col_name,
        labels.col_manufacturer,
        labels.col_model,
        labels.col_category,
        labels.col_serial,
        labels.col_checkout_date,
    ]
    tbl = doc.add_table(rows=1 + len(enriched), cols=len(columns))
    tbl.style = "Table Grid"

    # Header row
    for col_idx, col_name in enumerate(columns):
        cell = tbl.rows[0].cells[col_idx]
        style_header_cell(cell, col_name)
        if template == "uwagi":
            set_cell_bg(cell, "3d5a80")
            run = cell.paragraphs[0].runs[0]
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data rows
    for row_idx, item in enumerate(enriched, start=1):
        asset = item.asset
        date_str = (
            item.activity_date.strftime("%Y-%m-%d") if item.activity_date else
            (asset.last_checkout.strftime("%Y-%m-%d") if asset.last_checkout else "-")
        )
        values = [
            asset.asset_tag,
            asset.name or "-",
            asset.manufacturer_name,
            asset.model_name,
            asset.category_name,
            asset.serial_display,
            date_str,
        ]
        for col_idx, val in enumerate(values):
            style_body_cell(tbl.rows[row_idx].cells[col_idx], val)

    # Signature block
    # Left: the admin who issued (may differ per asset; use first or blank if mixed)
    admin_names = list({item.admin_name for item in enriched if item.admin_name})
    admin_display = admin_names[0] if len(admin_names) == 1 else ""

    add_signature_table(
        doc,
        left_label=labels.doc_issued_by,
        right_label=labels.doc_received_by,
        left_name=admin_display,
        right_name=user.display_name,
        date_label=labels.doc_date,
        sig_label=labels.doc_signature,
        name_label=labels.doc_name_surname,
    )

    if template == "uwagi":
        sig_tbl = doc.tables[-1]
        for cell in sig_tbl.rows[0].cells:
            set_cell_bg(cell, "3d5a80")
            if cell.paragraphs[0].runs:
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tmp = make_temp_file(".docx")
    doc.save(str(tmp))
    return tmp
