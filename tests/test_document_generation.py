"""Tests for DOCX document generation (no network calls needed)."""

from __future__ import annotations

from docx import Document

from app.documents.checkout import build_checkout_docx
from app.documents.return_form import build_return_docx
from app.i18n import get_labels
from app.snipeit.models import AssetWithActivity, User

EN = get_labels("en")
TR = get_labels("tr")


def test_checkout_docx_creates_file(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    path = build_checkout_docx([sample_enriched], sample_user, EN)
    assert path.exists()
    assert path.suffix == ".docx"
    path.unlink()


def test_checkout_docx_content(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    path = build_checkout_docx([sample_enriched], sample_user, EN)
    doc = Document(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert EN.doc_checkout_title in full_text
    assert sample_user.display_name in full_text
    path.unlink()


def test_checkout_docx_table_has_asset(
    sample_user: User, sample_enriched: AssetWithActivity
) -> None:
    path = build_checkout_docx([sample_enriched], sample_user, EN)
    doc = Document(str(path))
    # First table is the asset table; look for the asset tag in cells
    asset_table = doc.tables[0]
    cell_texts = [cell.text for row in asset_table.rows for cell in row.cells]
    assert sample_enriched.asset.asset_tag in cell_texts
    path.unlink()


def test_return_docx_creates_file(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    path = build_return_docx([sample_enriched], sample_user, EN)
    assert path.exists()
    assert path.suffix == ".docx"
    path.unlink()


def test_return_docx_content(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    path = build_return_docx([sample_enriched], sample_user, EN)
    doc = Document(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert EN.doc_return_title in full_text
    path.unlink()


def test_turkish_labels(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    path = build_checkout_docx([sample_enriched], sample_user, TR)
    doc = Document(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert TR.doc_checkout_title in full_text
    path.unlink()


def test_multiple_assets(sample_user: User, sample_enriched: AssetWithActivity) -> None:
    """Generating a form with multiple assets should produce one row per asset."""
    from copy import deepcopy


    second = deepcopy(sample_enriched)
    second.asset.id = 99
    second.asset.asset_tag = "00099"

    path = build_checkout_docx([sample_enriched, second], sample_user, EN)
    doc = Document(str(path))
    asset_table = doc.tables[0]
    # Header row + 2 data rows = 3 rows total
    assert len(asset_table.rows) == 3
    path.unlink()
